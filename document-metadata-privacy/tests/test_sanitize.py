"""Behavior tests: catch metadata leaks, content damage, and unsafe overwrites."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from pptx import Presentation
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, NameObject, DictionaryObject, ArrayObject, TextStringObject
from lxml import etree

SCRIPT = Path(__file__).resolve().parents[1] / 'scripts' / 'sanitize.py'


class SanitizeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=Path.cwd(), prefix='metadata-test-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def run_cli(self, *args, code=0):
        result = subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], capture_output=True, text=True)
        self.assertEqual(result.returncode, code, result.stdout + result.stderr)
        self.assertTrue(result.stdout.strip(), 'CLI must return a JSON result')
        return json.loads(result.stdout)

    def office(self, ext):
        path = self.root / ('source' + ext)
        if ext == '.docx':
            doc = Document()
            doc.add_paragraph('Keep Codex in the body: legitimate content.')
            doc.core_properties.author = 'Codex'
            doc.core_properties.last_modified_by = 'Private User'
            doc.save(path)
        elif ext == '.pptx':
            doc = Presentation()
            slide = doc.slides.add_slide(doc.slide_layouts[1])
            slide.shapes.title.text = 'Quarterly results'
            slide.notes_slide.notes_text_frame.text = 'Private speaker notes'
            doc.core_properties.author = 'Codex'
            doc.save(path)
        else:
            doc = Workbook()
            doc.active['A1'] = 12
            doc.active['A2'] = '=A1*2'
            doc.properties.creator = 'Codex'
            doc.save(path)
        with ZipFile(path) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        parts['docProps/custom.xml'] = b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"><property name="Generator" pid="2" fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"><vt:lpwstr xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">Codex</vt:lpwstr></property></Properties>'
        parts['docProps/app.xml'] = b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"><Application>WPS Office</Application><AppVersion>12.1</AppVersion><Company>Private Company</Company><Manager>Private User</Manager><Template>private.dotx</Template><TotalTime>50</TotalTime><Pages>1</Pages></Properties>'
        with ZipFile(path, 'w', ZIP_DEFLATED) as z:
            z.comment = b'Private User'
            for name, data in parts.items():
                z.writestr(name, data)
        return path, parts

    def test_office_properties_removed_and_other_parts_byte_identical(self):
        for ext in ('.docx', '.pptx', '.xlsx'):
            with self.subTest(ext=ext):
                source, before = self.office(ext)
                original = source.read_bytes()
                output = self.root / ('final' + ext)
                report = self.run_cli('clean', source, '--output', output)
                self.assertEqual(source.read_bytes(), original)
                self.assertEqual(report['remaining']['metadata_fields'], [])
                with ZipFile(output) as z:
                    self.assertEqual(z.comment, b'')
                    for name, data in before.items():
                        if name not in ('docProps/core.xml', 'docProps/custom.xml', 'docProps/app.xml'):
                            self.assertEqual(z.read(name), data, name)
                    for name in ('docProps/core.xml', 'docProps/custom.xml'):
                        self.assertEqual(len(etree.fromstring(z.read(name))), 0)
                    self.assertNotIn(b'WPS', z.read('docProps/app.xml'))
                if ext == '.docx':
                    self.assertIn('Keep Codex', Document(output).paragraphs[0].text)
                    self.assertTrue(report['remaining']['marker_parts'])
                elif ext == '.pptx':
                    self.assertEqual(Presentation(output).slides[0].shapes.title.text, 'Quarterly results')
                    self.assertIn('speaker_notes', report['remaining']['review_flags'])
                else:
                    self.assertEqual(load_workbook(output).active['A2'].value, '=A1*2')

    def test_inspect_does_not_modify_and_reports_field_names(self):
        source, _ = self.office('.docx')
        before = source.read_bytes()
        report = self.run_cli('inspect', source)
        self.assertTrue(report['metadata_fields'])
        self.assertEqual(source.read_bytes(), before)
        self.assertNotIn('Private User', json.dumps(report))

    def test_existing_output_and_in_place_rejected(self):
        source, _ = self.office('.docx')
        before = source.read_bytes()
        self.run_cli('clean', source, '--output', source, code=2)
        self.assertEqual(source.read_bytes(), before)
        out = self.root / 'existing.docx'
        out.write_bytes(b'keep me')
        self.run_cli('clean', source, '--output', out, code=2)
        self.assertEqual(out.read_bytes(), b'keep me')

    def test_office_signature_rejected_without_output(self):
        source, _ = self.office('.docx')
        with ZipFile(source, 'a') as z:
            z.writestr('_xmlsignatures/sig1.xml', '<Signature/>')
        output = self.root / 'signed-copy.docx'
        self.run_cli('clean', source, '--output', output, code=2)
        self.assertFalse(output.exists())

    def pdf(self, signed=False, encrypted=False, archival=False):
        source = self.root / 'source.pdf'
        writer = PdfWriter()
        page = writer.add_blank_page(300, 400)
        stream = DecodedStreamObject()
        stream.set_data(b'q 0.5 0.5 0.5 rg 10 10 50 50 re f Q')
        page[NameObject('/Contents')] = writer._add_object(stream)
        writer.add_metadata({'/Author': 'Codex', '/Producer': 'WPS Office', '/Private': 'Private User'})
        xmp = DecodedStreamObject()
        xmp.set_data(b'<x:xmpmeta xmlns:x="adobe:ns:meta/">Codex' + (b'pdfaid:' if archival else b'') + b'</x:xmpmeta>')
        xmp[NameObject('/Type')] = NameObject('/Metadata')
        xmp[NameObject('/Subtype')] = NameObject('/XML')
        ref = writer._add_object(xmp)
        writer.root_object[NameObject('/Metadata')] = ref
        page[NameObject('/Metadata')] = ref
        if signed:
            field = DictionaryObject({NameObject('/FT'): NameObject('/Sig'), NameObject('/T'): TextStringObject('Signature1')})
            writer.root_object[NameObject('/AcroForm')] = DictionaryObject({NameObject('/Fields'): ArrayObject([writer._add_object(field)])})
        if encrypted:
            writer.encrypt('secret')
        writer.write(source)
        return source

    def test_pdf_info_xmp_and_orphaned_metadata_removed_content_preserved(self):
        source = self.pdf()
        before = source.read_bytes()
        output = self.root / 'final.pdf'
        self.run_cli('clean', source, '--output', output)
        pdf = PdfReader(output)
        self.assertFalse(pdf.metadata)
        self.assertIsNone(pdf.xmp_metadata)
        self.assertNotIn('/Metadata', pdf.pages[0])
        self.assertEqual(len(pdf.pages), 1)
        self.assertEqual(pdf.pages[0].get_contents().get_data(), PdfReader(source).pages[0].get_contents().get_data())
        self.assertNotIn(b'Codex', output.read_bytes())
        self.assertNotIn(b'WPS Office', output.read_bytes())
        self.assertEqual(source.read_bytes(), before)

    def test_pdf_signature_encryption_and_archival_rejected(self):
        for option in ('signed', 'encrypted', 'archival'):
            with self.subTest(option=option):
                source = self.pdf(**{option: True})
                out = self.root / (option + '.pdf')
                self.run_cli('clean', source, '--output', out, code=2)
                self.assertFalse(out.exists())

    def test_legacy_wps_format_and_extension_mismatch_rejected(self):
        source = self.root / 'old.wps'
        source.write_bytes(b'legacy format')
        self.run_cli('clean', source, '--output', self.root / 'new.wps', code=2)
        source, _ = self.office('.docx')
        self.run_cli('clean', source, '--output', self.root / 'new.pdf', code=2)


if __name__ == '__main__':
    unittest.main()
